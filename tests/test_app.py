"""Tests for the FastAPI activities backend using AAA pattern."""

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_get_activities_returns_activity_data(client: TestClient):
    # Arrange
    expected_activity = "Chess Club"

    # Act
    response = client.get("/activities")

    # Assert
    assert response.status_code == 200
    activities = response.json()
    assert expected_activity in activities
    assert activities[expected_activity]["description"] == "Learn strategies and compete in chess tournaments"
    assert isinstance(activities[expected_activity]["participants"], list)


def test_signup_adds_participant_and_persists(client: TestClient, app_module_with_temp_data):
    # Arrange
    activity_name = "Chess Club"
    email = "newstudent@mergington.edu"

    # Act
    response = client.post(f"/activities/{activity_name}/signup", params={"email": email})

    # Assert
    assert response.status_code == 200
    assert email in app_module_with_temp_data.activities[activity_name]["participants"]

    workbook = load_workbook(app_module_with_temp_data.DATA_FILE)
    participants_sheet = workbook[app_module_with_temp_data.PARTICIPANTS_SHEET]
    persisted_emails = [row[1] for row in participants_sheet.iter_rows(min_row=2, values_only=True) if row[0] == activity_name]
    assert email in persisted_emails


def test_signup_duplicate_returns_400(client: TestClient):
    # Arrange
    activity_name = "Chess Club"
    email = "michael@mergington.edu"

    # Act
    response = client.post(f"/activities/{activity_name}/signup", params={"email": email})

    # Assert
    assert response.status_code == 400
    assert "already signed up" in response.json()["detail"]


def test_remove_participant_deletes_from_activity_and_persists(client: TestClient, app_module_with_temp_data):
    # Arrange
    activity_name = "Chess Club"
    email = app_module_with_temp_data.activities[activity_name]["participants"][0]

    # Act
    response = client.delete(f"/activities/{activity_name}/participants", params={"email": email})

    # Assert
    assert response.status_code == 200
    assert email not in app_module_with_temp_data.activities[activity_name]["participants"]

    workbook = load_workbook(app_module_with_temp_data.DATA_FILE)
    participants_sheet = workbook[app_module_with_temp_data.PARTICIPANTS_SHEET]
    persisted_emails = [row[1] for row in participants_sheet.iter_rows(min_row=2, values_only=True) if row[0] == activity_name]
    assert email not in persisted_emails


def test_create_activity_adds_new_activity(client: TestClient, app_module_with_temp_data):
    # Arrange
    new_activity = {
        "name": "Robotics Club",
        "description": "Build robots and compete",
        "schedule": "Saturdays, 10:00 AM - 12:00 PM",
        "max_participants": 20,
    }

    # Act
    response = client.post("/activities", json=new_activity)

    # Assert
    assert response.status_code == 200
    assert "Robotics Club" in app_module_with_temp_data.activities
    assert app_module_with_temp_data.activities["Robotics Club"]["description"] == new_activity["description"]

    workbook = load_workbook(app_module_with_temp_data.DATA_FILE)
    activities_sheet = workbook[app_module_with_temp_data.ACTIVITIES_SHEET]
    persisted_names = [row[0] for row in activities_sheet.iter_rows(min_row=2, values_only=True)]
    assert "Robotics Club" in persisted_names


def test_delete_activity_removes_activity(client: TestClient, app_module_with_temp_data):
    # Arrange
    activity_name = "Chess Club"

    # Act
    response = client.delete(f"/activities/{activity_name}")

    # Assert
    assert response.status_code == 200
    assert activity_name not in app_module_with_temp_data.activities

    workbook = load_workbook(app_module_with_temp_data.DATA_FILE)
    activities_sheet = workbook[app_module_with_temp_data.ACTIVITIES_SHEET]
    persisted_names = [row[0] for row in activities_sheet.iter_rows(min_row=2, values_only=True)]
    assert activity_name not in persisted_names
