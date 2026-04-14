"""
High School Management System API

A super simple FastAPI application that allows students to view and sign up
for extracurricular activities at Mergington High School.
"""

from fastapi import FastAPI, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.responses import RedirectResponse
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
import os
from pathlib import Path

app = FastAPI(title="Mergington High School API",
              description="API for viewing and signing up for extracurricular activities")

# Mount the static files directory
current_dir = Path(__file__).parent
app.mount("/static", StaticFiles(directory=os.path.join(Path(__file__).parent,
          "static")), name="static")

DATA_FILE = Path(__file__).parent / "activities.xlsx"
ACTIVITIES_SHEET = "activities"
PARTICIPANTS_SHEET = "participants"

INITIAL_ACTIVITIES = {
    "Chess Club": {
        "description": "Learn strategies and compete in chess tournaments",
        "schedule": "Fridays, 3:30 PM - 5:00 PM",
        "max_participants": 12,
        "participants": ["michael@mergington.edu", "daniel@mergington.edu"]
    },
    "Programming Class": {
        "description": "Learn programming fundamentals and build software projects",
        "schedule": "Tuesdays and Thursdays, 3:30 PM - 4:30 PM",
        "max_participants": 20,
        "participants": ["emma@mergington.edu", "sophia@mergington.edu"]
    },
    "Gym Class": {
        "description": "Physical education and sports activities",
        "schedule": "Mondays, Wednesdays, Fridays, 2:00 PM - 3:00 PM",
        "max_participants": 30,
        "participants": ["john@mergington.edu", "olivia@mergington.edu"]
    },
    "Basketball Team": {
        "description": "Competitive basketball league and practice drills",
        "schedule": "Tuesdays and Thursdays, 4:00 PM - 5:30 PM",
        "max_participants": 15,
        "participants": ["james@mergington.edu"]
    },
    "Tennis Club": {
        "description": "Tennis skills development and friendly matches",
        "schedule": "Mondays and Wednesdays, 3:30 PM - 4:30 PM",
        "max_participants": 16,
        "participants": ["sarah@mergington.edu", "alex@mergington.edu"]
    },
    "Art Painting": {
        "description": "Learn painting techniques and create original artwork",
        "schedule": "Wednesdays, 3:30 PM - 5:00 PM",
        "max_participants": 18,
        "participants": ["isabella@mergington.edu"]
    },
    "Drama Club": {
        "description": "Stage acting, theater production, and performance arts",
        "schedule": "Mondays and Fridays, 4:00 PM - 5:30 PM",
        "max_participants": 25,
        "participants": ["lucas@mergington.edu", "mia@mergington.edu"]
    },
    "Science Club": {
        "description": "Explore scientific experiments and research projects",
        "schedule": "Thursdays, 3:30 PM - 4:45 PM",
        "max_participants": 15,
        "participants": ["ava@mergington.edu", "noah@mergington.edu"]
    },
    "Debate Team": {
        "description": "Participate in academic debates and develop argumentation skills",
        "schedule": "Tuesdays, 3:30 PM - 4:45 PM",
        "max_participants": 12,
        "participants": ["ethan@mergington.edu"]
    }
}


class ActivityCreate(BaseModel):
    name: str
    description: str
    schedule: str
    max_participants: int


def create_workbook(path: Path, activities: dict):
    workbook = Workbook()
    activities_sheet = workbook.active
    activities_sheet.title = ACTIVITIES_SHEET
    activities_sheet.append(["activity_name", "description", "schedule", "max_participants"])

    participants_sheet = workbook.create_sheet(PARTICIPANTS_SHEET)
    participants_sheet.append(["activity_name", "email"])

    for activity_name, details in activities.items():
        activities_sheet.append([
            activity_name,
            details["description"],
            details["schedule"],
            details["max_participants"],
        ])
        for email in details["participants"]:
            participants_sheet.append([activity_name, email])

    workbook.save(path)


def load_data_from_excel(path: Path):
    if not path.exists():
        create_workbook(path, INITIAL_ACTIVITIES)

    workbook = load_workbook(path)
    activities = {}

    if ACTIVITIES_SHEET in workbook.sheetnames:
        sheet = workbook[ACTIVITIES_SHEET]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            activity_name, description, schedule, max_participants = row
            activities[activity_name] = {
                "description": description,
                "schedule": schedule,
                "max_participants": int(max_participants),
                "participants": [],
            }

    if PARTICIPANTS_SHEET in workbook.sheetnames:
        sheet = workbook[PARTICIPANTS_SHEET]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            activity_name, email = row
            if activity_name in activities:
                activities[activity_name]["participants"].append(email)

    return activities


def save_data_to_excel(path: Path, activities: dict):
    create_workbook(path, activities)


activities = load_data_from_excel(DATA_FILE)


@app.get("/")
def root():
    return RedirectResponse(url="/static/index.html")


@app.get("/activities")
def get_activities():
    return activities


@app.post("/activities/{activity_name}/signup")
def signup_for_activity(activity_name: str, email: str):        
    """Sign up a student for an activity"""
    # Validate activity exists
    if activity_name not in activities:
        raise HTTPException(status_code=404, detail="Activity not found")

    # Get the specific activity
    activity = activities[activity_name]

    # Validate student is not already signed up
    if email in activity["participants"]:
        raise HTTPException(status_code=400, detail="Student is already signed up for this activity")

    # Add student
    activity["participants"].append(email)
    save_data_to_excel(DATA_FILE, activities)
    return {"message": f"Signed up {email} for {activity_name}"}


@app.delete("/activities/{activity_name}/participants")
def remove_participant(activity_name: str, email: str):
    """Unregister a student from an activity"""
    if activity_name not in activities:
        raise HTTPException(status_code=404, detail="Activity not found")

    activity = activities[activity_name]

    if email not in activity["participants"]:
        raise HTTPException(status_code=404, detail="Participant not found")

    activity["participants"].remove(email)
    save_data_to_excel(DATA_FILE, activities)
    return {"message": f"Removed {email} from {activity_name}"}


@app.post("/activities")
def create_activity(activity: ActivityCreate):
    if activity.name in activities:
        raise HTTPException(status_code=400, detail="Activity already exists")

    activities[activity.name] = {
        "description": activity.description,
        "schedule": activity.schedule,
        "max_participants": activity.max_participants,
        "participants": [],
    }
    save_data_to_excel(DATA_FILE, activities)
    return {"message": f"Created activity {activity.name}"}


@app.delete("/activities/{activity_name}")
def delete_activity(activity_name: str):
    if activity_name not in activities:
        raise HTTPException(status_code=404, detail="Activity not found")

    activities.pop(activity_name)
    save_data_to_excel(DATA_FILE, activities)
    return {"message": f"Deleted activity {activity_name}"}
